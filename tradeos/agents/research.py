"""Research agent: reads reports you drop in and extracts sector / stock views with Claude,
cross-checked against price momentum. Re-runs whenever new or changed files appear."""
from __future__ import annotations

import glob
import json
from pathlib import Path

import pandas as pd

from ..analytics import rank_universe
from ..storage import dstr
from .base import BaseAgent, register

SYSTEM = """You are an equity research analyst for Indian markets (NSE/BSE).
You read reports (broker notes, results, macro/sector reports) and produce a structured, sober view.
Separate what the documents say from your inference. Flag low-confidence calls. Never promise returns.
Return ONLY a JSON object, no prose outside it."""

SCHEMA = """{
  "summary": "3-5 sentence overview",
  "sectors": [{"name": "...", "stance": "overweight|neutral|underweight", "confidence": 0.0-1.0,
               "reasons": ["..."], "momentum_agrees": true|false|null}],
  "stocks": [{"symbol": "NSE ticker e.g. RELIANCE", "view": "bullish|bearish|neutral",
              "reasons": ["..."], "source": "file name"}],
  "themes": ["..."],
  "risks": ["..."],
  "suggested_agents": [{"type": "alert|rule|rotation", "name": "...", "spec": {}, "why": "..."}]
}"""


def read_text(path: Path) -> str:
    suf = path.suffix.lower()
    try:
        if suf == ".pdf":
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        if suf == ".docx":
            import docx  # python-docx
            return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        if suf in (".txt", ".md", ".csv", ".json", ".html", ".htm"):
            return path.read_text(errors="ignore")
    except Exception as e:
        return f"[could not read {path.name}: {e}]"
    return f"[unsupported file type {suf}]"


@register
class ResearchAgent(BaseAgent):
    type_name = "research"
    description = ("Reads reports you drop into a folder (PDF, TXT, MD, CSV, DOCX: broker notes, results, "
                   "macro/sector reports) and uses Claude to extract which sectors/themes/stocks to focus on or avoid, "
                   "cross-checked with sector momentum. Suggests follow-up agents. Re-runs when new files arrive. "
                   "No trading, no backtest. Needs ANTHROPIC_API_KEY.")
    defaults = {"inputs": ["inbox/*"], "question": "Which sectors, themes and stocks should we focus on, and which "
                "should we avoid? What would change the view?", "cross_check_universe": "sectors",
                "max_chars": 150_000}
    example = {"inputs": ["inbox/q2fy27/*.pdf"],
               "question": "Which sectors show improving earnings momentum and where are downgrades clustering?"}

    @classmethod
    def _check(cls, s: dict) -> None:
        if isinstance(s["inputs"], str):
            s["inputs"] = [s["inputs"]]
        if not s["inputs"]:
            raise ValueError("research agent needs at least one input path/glob")

    def files(self) -> list[Path]:
        home, out = self.ctx.settings.home, set()
        for pat in self.spec["inputs"]:
            p = Path(pat).expanduser()
            pattern = str(p if p.is_absolute() else home / p)
            out.update(Path(f) for f in glob.glob(pattern, recursive=True) if Path(f).is_file())
        return sorted(out)

    def analyze(self, files: list[Path], asof=None) -> dict:
        per = max(3000, self.spec["max_chars"] // max(1, len(files)))
        docs = "\n\n".join(f"=== FILE: {f.name} ===\n{read_text(f)[:per]}" for f in files)
        quant = ""
        if self.spec.get("cross_check_universe"):
            try:
                r = rank_universe(self.ctx.provider, self.spec["cross_check_universe"], asof)
                quant = r.to_string(index=False)
            except Exception as e:
                quant = f"(momentum table unavailable: {e})"
        prompt = (f"Question: {self.spec['question']}\n\nPrice momentum table (score 0-1, higher = stronger):\n"
                  f"{quant}\n\nDocuments:\n{docs}\n\nReturn JSON with this shape:\n{SCHEMA}")
        return self.ctx.llm().complete_json(SYSTEM, prompt)

    def run_daily(self, asof) -> dict:
        files = self.files()
        seen = self.state.setdefault("seen", {})
        new = [f for f in files if seen.get(str(f)) != f.stat().st_mtime]
        if not new:
            return {"signals": [], "note": "no new reports"}
        res = self.analyze(files, asof)
        for f in files:
            seen[str(f)] = f.stat().st_mtime
        self.state["latest"] = res
        self.state["last_processed"] = dstr(pd.Timestamp(asof))
        self.save_state()
        self.ctx.store.add_run(self.id, "research", dstr(pd.Timestamp(asof)), res)
        ow = [s["name"] for s in res.get("sectors", []) if s.get("stance") == "overweight"]
        uw = [s["name"] for s in res.get("sectors", []) if s.get("stance") == "underweight"]
        msg = f"{len(new)} new report(s). Overweight: {', '.join(ow) or '-'} | Underweight: {', '.join(uw) or '-'}"
        self.signal(asof, "-", "RESEARCH", None, msg)
        return {"signals": self._emitted, "result": res}

    def backtest(self, start=None, end=None) -> dict:
        return {"note": "Research agents don't backtest. Create rule/rotation/alert agents from "
                        "'suggested_agents' in the latest result and backtest those.",
                "latest": self.state.get("latest")}

    def status(self) -> dict:
        latest = self.state.get("latest") or {}
        return {"files_tracked": len(self.state.get("seen", {})), "summary": latest.get("summary"),
                "sectors": latest.get("sectors", [])[:8], "suggested_agents": latest.get("suggested_agents", [])}

    def liquidate(self, asof, note: str = "") -> list:
        return []


def dumps(o) -> str:
    return json.dumps(o, indent=2, default=str)
